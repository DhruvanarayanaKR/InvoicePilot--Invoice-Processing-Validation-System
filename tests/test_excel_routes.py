from io import BytesIO

from openpyxl import Workbook, load_workbook

from invoicepilot import create_app


def _app(tmp_path):
    return create_app(
        {
            "TESTING": True,
            "APP_ENV": "testing",
            "SECRET_KEY": "test-key",
            "DATABASE_PATH": str(tmp_path / "excel-route-test.sqlite3"),
        }
    )


def test_excel_import_uses_validation_and_flags_bad_rows(tmp_path):
    app = _app(tmp_path)
    workbook = Workbook()
    sheet = workbook.active
    sheet.append(
        [
            "invoice_number",
            "vendor_name",
            "invoice_date",
            "due_date",
            "subtotal",
            "tax",
            "total",
        ]
    )
    sheet.append(
        [
            "XL-100",
            "Orchid Systems",
            "2026-09-01",
            "2026-09-30",
            100,
            8,
            115,
        ]
    )
    sheet.append(["XL-PARTIAL", None, None, None, None, None, None])
    stream = BytesIO()
    workbook.save(stream)
    stream.seek(0)

    client = app.test_client()
    response = client.post(
        "/excel/import",
        data={"excel_file": (stream, "records.xlsx")},
        content_type="multipart/form-data",
        follow_redirects=True,
    )

    assert response.status_code == 200
    assert b"XL-100" in response.data
    assert b"XL-PARTIAL" in response.data
    assert b"needs review" in response.data.lower()
    assert b"2 flagged for human review" in response.data


def test_excel_template_and_export_downloads_are_workbooks(tmp_path):
    client = _app(tmp_path).test_client()

    template_response = client.get("/excel/template")
    export_response = client.get("/excel/export")

    template = load_workbook(BytesIO(template_response.data), read_only=True)
    exported = load_workbook(BytesIO(export_response.data), read_only=True)
    assert template_response.status_code == 200
    assert template.active["A1"].value == "invoice_number"
    assert export_response.status_code == 200
    assert exported.active["A1"].value == "invoice_number"


def test_invalid_excel_template_is_reported_without_inserting_rows(tmp_path):
    app = _app(tmp_path)
    workbook = Workbook()
    workbook.active.append(["wrong_header"])
    stream = BytesIO()
    workbook.save(stream)
    stream.seek(0)

    response = app.test_client().post(
        "/excel/import",
        data={"excel_file": (stream, "wrong.xlsx")},
        content_type="multipart/form-data",
        follow_redirects=True,
    )

    assert response.status_code == 200
    assert b"Missing required Excel columns" in response.data
    assert b"0 total" in response.data
