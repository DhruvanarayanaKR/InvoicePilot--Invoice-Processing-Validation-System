from io import BytesIO

from openpyxl import Workbook, load_workbook

from invoicepilot.excel_service import export_invoices_to_excel, import_invoices_from_excel


def test_excel_export_contains_expected_columns_and_rows():
    invoices = [
        {
            "invoice_number": "EX-100",
            "vendor_name": "Brightline Co",
            "invoice_date": "2026-02-01",
            "due_date": "2026-02-28",
            "subtotal": 1200.0,
            "tax": 96.0,
            "total": 1296.0,
            "processing_status": "pending",
            "exception_reasons": ["None"],
            "source_filename": "invoice.pdf",
            "created_at": "2026-02-02 10:30:00",
        }
    ]

    workbook_stream = export_invoices_to_excel(invoices)
    workbook = load_workbook(workbook_stream)
    sheet = workbook.active

    assert sheet["A1"].value == "invoice_number"
    assert sheet["B2"].value == "Brightline Co"
    assert sheet["H2"].value == "pending"
    assert sheet.max_row == 2


def test_excel_export_treats_formula_like_invoice_number_as_text():
    workbook_stream = export_invoices_to_excel(
        [{"invoice_number": "=HYPERLINK(\"https://example.com\")"}]
    )
    worksheet = load_workbook(workbook_stream).active

    assert worksheet["A2"].value.startswith("'=")
    assert worksheet["A2"].data_type == "s"


def test_excel_import_reads_invoice_rows_from_workbook():
    workbook = Workbook()
    sheet = workbook.active
    sheet.append([
        "invoice_number",
        "vendor_name",
        "invoice_date",
        "due_date",
        "subtotal",
        "tax",
        "total",
        "processing_status",
        "exception_reasons",
        "source_filename",
        "created_at",
    ])
    sheet.append([
        "EX-200",
        "Harbor Labs",
        "2026-03-01",
        "2026-03-30",
        75.00,
        6.00,
        81.00,
        "needs_review",
        "Duplicate invoice",
        "harbor.xlsx",
        "2026-03-02 08:00:00",
    ])

    buffer = BytesIO()
    workbook.save(buffer)
    buffer.seek(0)

    records = import_invoices_from_excel(buffer)

    assert records[0]["invoice_number"] == "EX-200"
    assert records[0]["vendor_name"] == "Harbor Labs"
    assert records[0]["subtotal"] == "75.00"
    assert records[0]["total"] == "81.00"
