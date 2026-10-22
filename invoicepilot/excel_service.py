from datetime import date, datetime
from io import BytesIO

from openpyxl import Workbook, load_workbook


EXPORT_COLUMNS = [
    "invoice_number",
    "vendor_name",
    "invoice_date",
    "due_date",
    "subtotal",
    "tax",
    "total",
    "processing_status",
    "exception_reasons",
    "source_filename",
    "created_at",
]
REQUIRED_IMPORT_COLUMNS = {
    "invoice_number",
    "vendor_name",
    "invoice_date",
    "due_date",
    "subtotal",
    "tax",
    "total",
}
AMOUNT_COLUMNS = {"subtotal", "tax", "total"}


def export_invoices_to_excel(invoices):
    """Return an Excel workbook as a BytesIO stream for download."""
    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = "Invoices"
    worksheet.append(EXPORT_COLUMNS)

    for invoice in invoices:
        row = [
            invoice.get("invoice_number"),
            invoice.get("vendor_name"),
            invoice.get("invoice_date"),
            invoice.get("due_date"),
            invoice.get("subtotal"),
            invoice.get("tax"),
            invoice.get("total"),
            invoice.get("processing_status"),
            ", ".join(invoice.get("exception_reasons", [])),
            invoice.get("source_filename"),
            invoice.get("created_at"),
        ]
        worksheet.append([_safe_excel_value(value) for value in row])

    return _save_workbook(workbook)


def create_import_template():
    """Return a blank workbook using the importer's canonical column names."""
    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = "Invoices"
    worksheet.append(EXPORT_COLUMNS)
    return _save_workbook(workbook)


def import_invoices_from_excel(file_obj):
    """Read invoice rows and normalize Excel dates and numeric cells."""
    workbook = load_workbook(file_obj, read_only=True, data_only=True)
    try:
        worksheet = workbook.active
        rows = worksheet.iter_rows(values_only=True)
        header_values = next(rows, None)
        if not header_values:
            return []

        headers = [
            str(value).strip().lower() if value is not None else ""
            for value in header_values
        ]
        if len(headers) != len(set(header for header in headers if header)):
            raise ValueError("The Excel template contains duplicate column headings.")

        missing_headers = REQUIRED_IMPORT_COLUMNS - set(headers)
        if missing_headers:
            missing_list = ", ".join(sorted(missing_headers))
            raise ValueError(f"Missing required Excel columns: {missing_list}.")

        records = []
        for row in rows:
            if not any(value is not None and str(value).strip() for value in row):
                continue

            record = {}
            for index, header in enumerate(headers):
                if not header or index >= len(row):
                    continue
                record[header] = _normalize_cell(header, row[index])
            records.append(record)

        return records
    finally:
        workbook.close()


def _normalize_cell(header, value):
    if value is None:
        return ""
    if header in {"invoice_date", "due_date"}:
        if isinstance(value, datetime):
            return value.date().isoformat()
        if isinstance(value, date):
            return value.isoformat()
    if header in AMOUNT_COLUMNS:
        if isinstance(value, (int, float)):
            return f"{value:.2f}"
    if isinstance(value, str):
        return value.strip()
    return str(value).strip()


def _safe_excel_value(value):
    if isinstance(value, str) and value.startswith(("=", "+", "-", "@")):
        return "'" + value
    return value


def _save_workbook(workbook):
    stream = BytesIO()
    workbook.save(stream)
    stream.seek(0)
    return stream
